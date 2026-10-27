"""Extract complete text before asking a model to interpret its fields."""
from pathlib import Path
import pdfplumber
from docx import Document
from PIL import Image
import pytesseract
from agent.config import rules

def harvest(path: str):
    path = Path(path)
    if path.stat().st_size > rules()["max_upload_mb"] * 1024 * 1024:
        raise ValueError("Invoice exceeds configured size limit")
    if path.suffix.lower() == ".docx":
        doc = Document(path)
        lines = [p.text for p in doc.paragraphs]
        lines += [" | ".join(c.text for c in r.cells) for t in doc.tables for r in t.rows]
        return {"text": "\n".join(lines), "method": "docx"}
    if path.suffix.lower() == ".png":
        with Image.open(path) as image:
            return {"text": pytesseract.image_to_string(image, lang="eng+spa+deu"), "method": "tesseract"}
    if path.suffix.lower() == ".pdf":
        texts, methods = [], []
        with pdfplumber.open(path) as pdf:
            if len(pdf.pages) > 100: raise ValueError("Maximum 100 invoice pages")
            for page in pdf.pages:
                text = page.extract_text() or ""
                method = "pdfplumber"
                if not text.strip():
                    text = pytesseract.image_to_string(page.to_image(resolution=200).original, lang="eng+spa+deu")
                    method = "tesseract"
                texts.append(text); methods.append(method)
        return {"text": "\n".join(texts), "method": "+".join(sorted(set(methods)))}
    raise ValueError("Supported invoice formats: PDF, DOCX, PNG")
