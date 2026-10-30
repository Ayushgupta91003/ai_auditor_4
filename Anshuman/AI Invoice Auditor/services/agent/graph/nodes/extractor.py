"""
Extractor Node
Handles file ingestion, checksum calculation, and text extraction from PDF, DOCX, and PNG formats.
Addresses requirement 3.1 (Monitoring & Checksum) and 3.2 (Extraction & OCR Fallback).
Digital PDFs use pdfplumber/pypdf; scanned PDFs and images fall back to Tesseract OCR
(English, Spanish and German language packs). No invoice content is hard-coded.
"""

import os
import json
import hashlib
import logging
from pathlib import Path
from typing import Dict, Any, Tuple
from ..state import InvoiceState
from ...utils.db import audit_db

logger = logging.getLogger("extractor_node")

OCR_LANGS = os.getenv("OCR_LANGS", "eng+spa+deu")

def compute_checksum(filepath: str) -> str:
    hasher = hashlib.sha256()
    with open(filepath, "rb") as f:
        while chunk := f.read(65536):
            hasher.update(chunk)
    return hasher.hexdigest()

def load_metadata(filepath: str) -> Dict[str, Any]:
    p = Path(filepath)
    meta_path = p.with_name(p.stem + ".meta.json")
    if meta_path.exists():
        try:
            with open(meta_path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            logger.warning(f"Could not read metadata for {filepath}: {e}")
    return {}

def ocr_image(filepath: str) -> str:
    """Runs Tesseract OCR on an image file."""
    try:
        import pytesseract
        from PIL import Image
        with Image.open(filepath) as img:
            text = pytesseract.image_to_string(img, lang=OCR_LANGS)
        return text.strip()
    except Exception as e:
        logger.error(f"OCR failed for image {filepath}: {e}")
        return ""

def ocr_pdf(filepath: str) -> str:
    """Renders PDF pages to images (Poppler) and runs Tesseract OCR on each."""
    try:
        import pytesseract
        from pdf2image import convert_from_path
        pages = convert_from_path(filepath, dpi=300)
        texts = []
        for page in pages:
            texts.append(pytesseract.image_to_string(page, lang=OCR_LANGS))
        return "\n".join(texts).strip()
    except Exception as e:
        logger.error(f"OCR failed for PDF {filepath}: {e}")
        return ""

def extract_pdf_text(filepath: str) -> Tuple[str, str]:
    text = ""
    method = "pdfplumber"
    try:
        import pdfplumber
        with pdfplumber.open(filepath) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
            text = "\n".join(pages).strip()
    except Exception as e:
        logger.warning(f"pdfplumber failed for {filepath}: {e}. Trying pypdf fallback.")
        try:
            import pypdf
            reader = pypdf.PdfReader(filepath)
            pages = [page.extract_text() or "" for page in reader.pages]
            text = "\n".join(pages).strip()
            method = "pypdf"
        except Exception as ex:
            logger.error(f"pypdf also failed: {ex}")

    # Fallback to OCR if embedded text layer is absent or empty
    if not text or len(text) < 15:
        logger.info(f"Direct text layer insufficient for {filepath}. Attempting OCR fallback.")
        ocr_text = ocr_pdf(filepath)
        if ocr_text:
            return ocr_text, "tesseract_ocr"

    return text, method

def extract_docx_text(filepath: str) -> Tuple[str, str]:
    import docx
    doc = docx.Document(filepath)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    # Extract table rows as well
    table_lines = []
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            table_lines.append(" | ".join(cells))

    all_text = "\n".join(paragraphs + table_lines).strip()
    return all_text, "python-docx"

def extractor_node(state: InvoiceState) -> InvoiceState:
    file_path = state["file_path"]
    logger.info(f"Extractor processing: {file_path}")

    # 1. Compute Checksum
    checksum = compute_checksum(file_path)
    state["file_checksum"] = checksum

    # 2. Check if already audited
    if audit_db.is_checksum_audited(checksum):
        logger.info(f"Invoice with checksum {checksum[:12]}... has already been audited. Skipping.")
        state["skipped"] = True
        state["skip_reason"] = "Already present in audit trail (checksum match)"
        return state

    state["skipped"] = False

    # 3. Load Metadata
    state["metadata"] = load_metadata(file_path)

    # 4. Extract Text
    ext = Path(file_path).suffix.lower()
    raw_text = ""
    method = "unknown"

    if ext == ".pdf":
        raw_text, method = extract_pdf_text(file_path)
    elif ext in [".docx", ".doc"]:
        raw_text, method = extract_docx_text(file_path)
    elif ext in [".png", ".jpg", ".jpeg", ".tiff"]:
        raw_text = ocr_image(file_path)
        method = "tesseract_ocr"
    else:
        logger.error(f"Unsupported file format: {ext}")
        state["error"] = f"Unsupported file format: {ext}"
        return state

    state["raw_extracted_text"] = raw_text
    state["extraction_method"] = method

    if not raw_text:
        state["error"] = f"No text could be extracted from {Path(file_path).name} (method: {method})."
        logger.warning(state["error"])
    else:
        logger.info(f"Extracted {len(raw_text)} chars from {file_path} via {method}")
    return state
